import os
import shutil
import platform
import subprocess
import tempfile
import hashlib
import zipfile
import tarfile
import re
import json
import time
from pathlib import Path
from datetime import datetime
from collections import defaultdict

try:
    import send2trash
    _SEND2TRASH = True
except ImportError:
    _SEND2TRASH = False

try:
    import win32api
    import win32file
    _WIN32_AVAILABLE = True
except ImportError:
    _WIN32_AVAILABLE = False

from core.undo import push_undo

_OS = platform.system()

# Undo keeps a file's previous contents in memory so `write` can be reversed.
_UNDO_CONTENT_LIMIT = 1_000_000


# ── REAL WINDOWS FOLDER LOCATIONS (handles OneDrive-redirected folders) ──────

def _win_known_folder(guid_str: str):
    """Ask Windows where a special folder REALLY is (Desktop may live in OneDrive)."""
    try:
        import ctypes
        from ctypes import wintypes

        class GUID(ctypes.Structure):
            _fields_ = [("Data1", wintypes.DWORD), ("Data2", wintypes.WORD),
                        ("Data3", wintypes.WORD), ("Data4", ctypes.c_ubyte * 8)]

        g = GUID()
        ctypes.windll.ole32.CLSIDFromString(guid_str, ctypes.byref(g))
        ptr = ctypes.c_wchar_p()
        if ctypes.windll.shell32.SHGetKnownFolderPath(
                ctypes.byref(g), 0, None, ctypes.byref(ptr)) == 0:
            p = Path(ptr.value)
            ctypes.windll.ole32.CoTaskMemFree(ptr)
            return p
    except Exception:
        pass
    return None


_KF_DESKTOP   = "{B4BFCC3A-DB2C-424C-B029-7FE99A87C641}"
_KF_DOWNLOADS = "{374DE290-123F-4565-9164-39C4925E467B}"
_KF_DOCUMENTS = "{FDD39AD0-238F-46AF-ADB4-6C85480369C7}"
_KF_PICTURES  = "{33E28130-4E1E-4676-835A-98395C3BC3BB}"
_KF_MUSIC     = "{4BD8D571-6D19-48D3-BE97-422220080E43}"
_KF_VIDEOS    = "{18989B1D-99B5-455B-841C-AB7C74E4DDFC}"


def _known_dir(guid: str, xdg_var: str, fallback_name: str) -> Path:
    if _OS == "Windows":
        p = _win_known_folder(guid)
        if p:
            return p
    if _OS == "Linux":
        xdg = os.environ.get(xdg_var, "")
        if xdg and Path(xdg).exists():
            return Path(xdg)
    return Path.home() / fallback_name


def _get_desktop() -> Path:
    return _known_dir(_KF_DESKTOP, "XDG_DESKTOP_DIR", "Desktop")


def _get_downloads() -> Path:
    return _known_dir(_KF_DOWNLOADS, "XDG_DOWNLOAD_DIR", "Downloads")


def _get_documents() -> Path:
    return _known_dir(_KF_DOCUMENTS, "XDG_DOCUMENTS_DIR", "Documents")


def _get_pictures() -> Path:
    return _known_dir(_KF_PICTURES, "XDG_PICTURES_DIR", "Pictures")


def _get_music() -> Path:
    return _known_dir(_KF_MUSIC, "XDG_MUSIC_DIR", "Music")


def _get_videos() -> Path:
    return _known_dir(_KF_VIDEOS, "XDG_VIDEOS_DIR", "Videos")


def _get_temp() -> Path:
    return Path(tempfile.gettempdir())


def _get_appdata() -> Path:
    app_data = os.environ.get("APPDATA")
    return Path(app_data) if app_data else Path.home()


def _get_onedrive() -> Path | None:
    one_drive = os.environ.get("OneDrive") or os.environ.get("OneDriveConsumer")
    if one_drive and Path(one_drive).exists():
        return Path(one_drive)
    p = Path.home() / "OneDrive"
    return p if p.exists() else None


# ── Stateful directory tracking & history ────────────────────────────────────
_CURRENT_DIR: Path = _get_desktop()
_DIR_HISTORY_BACK: list[Path] = []
_DIR_HISTORY_FORWARD: list[Path] = []

# The ONE File Explorer window this tool controls. Opening folders re-uses it
# (like clicking through folders) unless the user asks for a NEW window.
_EXPLORER_HWND: int | None = None


def get_current_dir() -> Path:
    return _CURRENT_DIR


def set_current_dir(p: Path):
    global _CURRENT_DIR
    _push_navigation(p)
    _CURRENT_DIR = p


def _push_navigation(new_dir: Path):
    global _CURRENT_DIR, _DIR_HISTORY_BACK, _DIR_HISTORY_FORWARD
    try:
        resolved = new_dir.resolve()
        if _CURRENT_DIR and resolved != _CURRENT_DIR.resolve():
            _DIR_HISTORY_BACK.append(_CURRENT_DIR)
            if len(_DIR_HISTORY_BACK) > 50:
                _DIR_HISTORY_BACK.pop(0)
            _DIR_HISTORY_FORWARD.clear()
    except Exception:
        pass


def _need(msg: str) -> str:
    """Signal to the AI model that it must ask the user, not guess."""
    return f"NEED_INFO: {msg} Ask the user this question and wait for the answer. Do NOT guess."


# ── Undo helpers ─────────────────────────────────────────────────────────────

def _undo_move(src: Path, dst: Path):
    """Reverse of a move: put it back where it came from."""
    def _fn():
        if not dst.exists():
            return f"'{dst.name}' is no longer there — nothing moved back."
        src.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(dst), str(src))
        return f"'{src.name}' is back in {src.parent.name}/."
    return _fn


def _undo_create(target: Path):
    """Reverse of a create: remove what we made — and only if we still made it."""
    def _fn():
        if not target.exists():
            return f"'{target.name}' is already gone."
        if target.is_dir():
            if any(target.iterdir()):
                return (f"'{target.name}' is not empty any more — "
                        f"leaving it alone rather than deleting your files.")
            target.rmdir()
        else:
            target.unlink()
        return f"Removed '{target.name}'."
    return _fn


def _undo_write(target: Path, previous: str | None):
    """Reverse of a write: restore old contents or remove newly created file."""
    def _fn():
        if previous is None:
            if target.exists():
                target.unlink()
                return f"Removed '{target.name}' — it did not exist before."
            return f"'{target.name}' is already gone."
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(previous, encoding="utf-8")
        return f"Restored the previous contents of '{target.name}'."
    return _fn


def _restore_from_trash(original: Path) -> str:
    """Best-effort undelete from Recycle Bin."""
    if _OS == "Windows":
        try:
            import win32com.client
            shell = win32com.client.Dispatch("Shell.Application")
            bin_folder = shell.NameSpace(10)      # ssfBITBUCKET
            for item in bin_folder.Items():
                if str(bin_folder.GetDetailsOf(item, 1)).strip().lower() == \
                        str(original.parent).strip().lower():
                    if str(item.Name).strip().lower() == original.name.strip().lower():
                        item.InvokeVerb("UNDELETE")
                        return f"'{original.name}' restored from the Recycle Bin."
        except Exception as e:
            print(f"[file] Recycle Bin restore failed: {e}")
    return (f"'{original.name}' is in the Recycle Bin — I could not pull it back "
            f"automatically, but it is there and can be restored by hand.")


def _is_safe_path(target: Path) -> bool:
    """Validate path permissions:
    - ALLOWED: all non-C: drives, user folders on C:
    - BLOCKED: Windows OS folders on C: (Windows, Program Files)
    """
    try:
        resolved = target.resolve()
        if _OS == "Windows":
            drive = resolved.drive.upper()
            if drive and drive != "C:":
                return True
            sys_roots = [
                Path("C:/Windows").resolve(),
                Path("C:/Program Files").resolve(),
                Path("C:/Program Files (x86)").resolve(),
            ]
            for s_root in sys_roots:
                if resolved == s_root or resolved.is_relative_to(s_root):
                    return False
            return True
        return True
    except Exception:
        return True


# ── Path resolution ──────────────────────────────────────────────────────────

def _drive_letter(expr: str) -> str | None:
    """'d:' / 'd drive' / 'drive d' / 'drive d:' -> 'D'. Plain single letters do NOT match."""
    m = re.match(r"^(?:in\s+)?(?:drive\s+([a-zA-Z])|([a-zA-Z])(?:\s+drive|:))\s*:?$", expr.strip().lower())
    if m:
        return (m.group(1) or m.group(2)).upper()
    return None


def _resolve_path(raw) -> Path:
    """Conversational path resolution with multi-drive, navigation and alias support.
    Accepts str or Path."""
    raw = str(raw or "").strip().strip('"').strip("'")

    shortcuts: dict[str, Path] = {
        "desktop":       _get_desktop(),
        "on desktop":    _get_desktop(),
        "downloads":     _get_downloads(),
        "in downloads":  _get_downloads(),
        "documents":     _get_documents(),
        "in documents":  _get_documents(),
        "pictures":      _get_pictures(),
        "music":         _get_music(),
        "videos":        _get_videos(),
        "temp":          _get_temp(),
        "tmp":           _get_temp(),
        "appdata":       _get_appdata(),
        "home":          Path.home(),
        "here":          _CURRENT_DIR,
        "current":       _CURRENT_DIR,
        "current folder": _CURRENT_DIR,
        "current directory": _CURRENT_DIR,
        ".":             _CURRENT_DIR,
        "":              _CURRENT_DIR,
        "up":            _CURRENT_DIR.parent,
        "parent":        _CURRENT_DIR.parent,
        "..":            _CURRENT_DIR.parent,
    }

    od = _get_onedrive()
    if od:
        shortcuts["onedrive"] = od
        shortcuts["one drive"] = od

    norm = raw.replace("\\", "/")
    lower = norm.lower()

    if lower in shortcuts:
        return shortcuts[lower]

    letter = _drive_letter(lower.rstrip("/"))
    if letter:
        return Path(f"{letter}:/")

    # "drive d/projects", "desktop/folder", "D:/stuff"
    head, sep, rest = norm.partition("/")
    if sep and head.strip():
        head_clean = head.strip().lower()
        letter = _drive_letter(head_clean)
        if letter:
            return Path(f"{letter}:/") / rest.strip("/")
        if head_clean in shortcuts:
            base = shortcuts[head_clean]
            return base / rest.strip("/") if rest.strip("/") else base

    # "D:test.txt" -> "D:/test.txt"
    if _OS == "Windows" and len(raw) >= 2 and raw[1] == ":" and (len(raw) == 2 or raw[2] not in ("/", "\\")):
        raw = raw[:2] + "/" + raw[2:]

    p = Path(raw).expanduser()
    if p.is_absolute():
        return p

    # Relative: look in current dir, then common places
    candidates = [_CURRENT_DIR, _get_desktop(), _get_downloads(), _get_documents()]
    if _OS == "Windows":
        candidates.append(Path("D:/"))
    for base in candidates:
        try:
            if (base / raw).exists():
                return base / raw
        except Exception:
            continue

    # Case-insensitive match in current dir
    try:
        if _CURRENT_DIR.exists() and _CURRENT_DIR.is_dir():
            for item in _CURRENT_DIR.iterdir():
                if item.name.lower() == lower:
                    return item
    except Exception:
        pass

    return _CURRENT_DIR / raw


def _format_size(b) -> str:
    b = float(b)
    for unit in ["B", "KB", "MB", "GB", "TB"]:
        if b < 1024:
            return f"{b:.1f} {unit}"
        b /= 1024
    return f"{b:.1f} PB"


def _safe_trash(target: Path) -> str:
    if not _SEND2TRASH:
        return (
            "send2trash is not installed. "
            "Permanent deletion is disabled for safety."
        )
    send2trash.send2trash(str(target))
    return f"Moved to Trash: {target.name}"


# ── Search across all locations ──────────────────────────────────────────────

_SKIP_DIRS = {"windows", "program files", "program files (x86)", "node_modules",
              "appdata", "$recycle.bin", "system volume information", ".git"}


def _search_roots() -> list[Path]:
    roots = [_CURRENT_DIR, _get_desktop(), _get_downloads(), _get_documents()]
    if _OS == "Windows":
        for letter in "DEFGH":
            if Path(f"{letter}:/").exists():
                roots.append(Path(f"{letter}:/"))
    out, seen = [], set()
    for r in roots:
        try:
            key = str(r.resolve()).lower()
        except Exception:
            key = str(r).lower()
        if key not in seen and r.exists():
            seen.add(key)
            out.append(r)
    return out


def _search_everywhere(name: str = "", extension: str = "", limit: int = 10,
                       time_cap: float = 12.0, roots: list[Path] | None = None) -> list[Path]:
    """Search Desktop, Downloads, Documents and every other drive for matching files."""
    roots = roots or _search_roots()
    ext = ("." + extension.lstrip(".")).lower() if extension else ""
    q = name.lower()
    seen, found = set(), []
    start = time.time()
    for root in roots:
        try:
            for dp, dns, fns in os.walk(root):
                dns[:] = [d for d in dns
                          if not d.startswith((".", "$")) and d.lower() not in _SKIP_DIRS]
                for fn in fns:
                    if q and q not in fn.lower():
                        continue
                    if ext and not fn.lower().endswith(ext):
                        continue
                    p = Path(dp) / fn
                    key = str(p).lower()
                    if key not in seen:
                        seen.add(key)
                        found.append(p)
                        if len(found) >= limit:
                            return found
                if time.time() - start > time_cap:
                    return found
        except Exception:
            continue
    return found


# ── WINDOWS FILE EXPLORER INTEGRATION (re-uses ONE window) ───────────────────

def _explorer_windows() -> list:
    """COM objects of every open File Explorer folder window (Windows only)."""
    import win32com.client
    shell = win32com.client.Dispatch("Shell.Application")
    wins = []
    for w in shell.Windows():
        try:
            if os.path.basename(str(w.FullName)).lower() == "explorer.exe":
                wins.append(w)
        except Exception:
            continue
    return wins


def _explorer_hwnds() -> set:
    try:
        return {int(w.HWND) for w in _explorer_windows()}
    except Exception:
        return set()


def _find_tracked_window():
    """The Explorer window this tool opened earlier, if it is still open — else None."""
    global _EXPLORER_HWND
    if _EXPLORER_HWND is None:
        return None
    try:
        for w in _explorer_windows():
            if int(w.HWND) == _EXPLORER_HWND:
                return w
    except Exception:
        pass
    _EXPLORER_HWND = None
    return None


def _bring_to_front(hwnd: int):
    try:
        import ctypes
        user32 = ctypes.windll.user32
        user32.ShowWindow(hwnd, 9)          # SW_RESTORE
        user32.SetForegroundWindow(hwnd)
    except Exception:
        pass


def _launch_new_explorer(args: list) -> None:
    """Start a NEW Explorer window and remember it as the one to re-use."""
    global _EXPLORER_HWND
    before = _explorer_hwnds()
    subprocess.Popen(args)
    for _ in range(40):                     # wait up to ~4 s for the window to appear
        time.sleep(0.1)
        new = _explorer_hwnds() - before
        if new:
            _EXPLORER_HWND = sorted(new)[-1]
            _bring_to_front(_EXPLORER_HWND)
            return


def _open_in_explorer(target: Path, new_window: bool = False) -> str:
    """Show a folder in File Explorer.
    Re-uses the window opened earlier (just navigates it) unless new_window=True.
    Returns a short status phrase."""
    global _EXPLORER_HWND
    if _OS == "Darwin":
        subprocess.Popen(["open", str(target)])
        return "Opened in Finder"
    if _OS != "Windows":
        subprocess.Popen(["xdg-open", str(target)])
        return "Opened in the file manager"

    try:
        if not new_window:
            w = _find_tracked_window()
            if w is not None:
                try:
                    w.Navigate2(str(target))
                    _bring_to_front(int(w.HWND))
                    return "Opened in the SAME File Explorer window"
                except Exception:
                    _EXPLORER_HWND = None       # window died mid-way → open a fresh one
        _launch_new_explorer(["explorer.exe", str(target)])
        return "Opened in a NEW File Explorer window"
    except ImportError:
        os.startfile(str(target))
        return "Opened in File Explorer (cannot re-use the window — install pywin32: pip install pywin32)"


def open_folder(path: str = "", new_window: bool = False) -> str:
    """Open a folder in File Explorer (same window by default), track it as current, and list contents."""
    try:
        global _CURRENT_DIR
        target = _resolve_path(path) if path else _CURRENT_DIR
        if target.is_file():
            target = target.parent
        if not target.exists():
            return _need(f"The folder '{target}' does not exist. Which folder do you mean?")
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        _push_navigation(target)
        _CURRENT_DIR = target

        status = _open_in_explorer(target, new_window=new_window)
        contents = list_files(str(target))
        return f"{status}: {target.resolve()}\n\n{contents}"
    except Exception as e:
        return f"Error opening folder: {e}"


def reveal_in_explorer(path: str = "", name: str = "", new_window: bool = False) -> str:
    """Show a file/folder highlighted in File Explorer (same window by default)."""
    global _EXPLORER_HWND
    try:
        base = _resolve_path(path) if path else _CURRENT_DIR
        target = (base / name) if (name and not base.name.lower() == name.lower()) else base
        if not target.exists():
            return _need(f"I could not find '{target}' to reveal. Which folder is it in?")
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        resolved_str = str(target.resolve())
        if _OS == "Windows":
            try:
                if not new_window:
                    w = _find_tracked_window()
                    if w is not None:
                        try:
                            w.Navigate2(str(target.parent))
                            time.sleep(0.3)
                            for _ in range(40):
                                if not w.Busy:
                                    break
                                time.sleep(0.1)
                            doc = w.Document
                            item = doc.Folder.ParseName(target.name)
                            if item is None:
                                raise RuntimeError("item not found in window")
                            doc.SelectItem(item, 29)     # select + deselect others + scroll into view + focus
                            _bring_to_front(int(w.HWND))
                            return f"Selected '{target.name}' in the SAME File Explorer window."
                        except Exception:
                            _EXPLORER_HWND = None
                _launch_new_explorer(["explorer.exe", "/select,", resolved_str])
                return f"Revealed and selected '{target.name}' in a NEW File Explorer window."
            except ImportError:
                subprocess.Popen(["explorer.exe", "/select,", resolved_str])
                return f"Revealed '{target.name}' in File Explorer (install pywin32 to re-use one window)."
        elif _OS == "Darwin":
            subprocess.Popen(["open", "-R", resolved_str])
            return f"Revealed '{target.name}' in Finder."
        else:
            subprocess.Popen(["xdg-open", str(target.parent)])
            return f"Opened folder containing '{target.name}' in file manager."
    except Exception as e:
        return f"Could not reveal item in File Explorer: {e}"


def get_open_explorers() -> str:
    """List all folders currently open in File Explorer windows."""
    if _OS != "Windows":
        return "Open explorer window listing is only available on Windows."
    try:
        import win32com.client
        import urllib.parse
        shell = win32com.client.Dispatch("Shell.Application")
        open_paths = []
        for win in shell.Windows():
            try:
                url = getattr(win, "LocationURL", "")
                if url.startswith("file:///"):
                    parsed = urllib.parse.unquote(url[8:].replace("/", "\\"))
                    if parsed and parsed not in open_paths:
                        open_paths.append(parsed)
            except Exception:
                continue

        if not open_paths:
            return "No File Explorer folder windows are currently detected."

        lines = [f"Currently open File Explorer windows ({len(open_paths)}):"]
        for p in open_paths:
            lines.append(f"  📁 {p}")
        return "\n".join(lines)
    except Exception as e:
        return f"Could not inspect open File Explorer windows: {e}"


def close_explorer(path: str = "") -> str:
    """Close specific or all open File Explorer folder windows."""
    if _OS != "Windows":
        return "Closing explorer windows is only supported on Windows."
    try:
        import win32com.client
        import urllib.parse
        shell = win32com.client.Dispatch("Shell.Application")
        closed_count = 0
        target_clean = path.lower().strip().replace("/", "\\") if path else ""

        for win in list(shell.Windows()):
            try:
                url = getattr(win, "LocationURL", "")
                if url.startswith("file:///"):
                    parsed = urllib.parse.unquote(url[8:].replace("/", "\\")).lower()
                    if not target_clean or target_clean in parsed:
                        win.Quit()
                        closed_count += 1
            except Exception:
                continue

        if closed_count > 0:
            return f"Closed {closed_count} File Explorer window(s)."
        return "No matching File Explorer windows were found to close."
    except Exception as e:
        return f"Could not close File Explorer window: {e}"


# ── NAVIGATION STACK ─────────────────────────────────────────────────────────

def navigate_back() -> str:
    global _CURRENT_DIR, _DIR_HISTORY_BACK, _DIR_HISTORY_FORWARD
    if not _DIR_HISTORY_BACK:
        return f"No previous folder in history. Currently at: {_CURRENT_DIR}"
    prev = _DIR_HISTORY_BACK.pop()
    _DIR_HISTORY_FORWARD.append(_CURRENT_DIR)
    _CURRENT_DIR = prev
    try:
        _open_in_explorer(prev)
    except Exception:
        pass
    contents = list_files(str(prev))
    return f"Navigated back to {prev.resolve()}.\n\n{contents}"


def navigate_forward() -> str:
    global _CURRENT_DIR, _DIR_HISTORY_BACK, _DIR_HISTORY_FORWARD
    if not _DIR_HISTORY_FORWARD:
        return f"No forward folder in history. Currently at: {_CURRENT_DIR}"
    next_dir = _DIR_HISTORY_FORWARD.pop()
    _DIR_HISTORY_BACK.append(_CURRENT_DIR)
    _CURRENT_DIR = next_dir
    try:
        _open_in_explorer(next_dir)
    except Exception:
        pass
    contents = list_files(str(next_dir))
    return f"Navigated forward to {next_dir.resolve()}.\n\n{contents}"


def navigate_up() -> str:
    parent = _CURRENT_DIR.parent
    if parent == _CURRENT_DIR:
        return f"Already at the root: {_CURRENT_DIR}"
    return open_folder(str(parent))


def get_navigation_history() -> str:
    crumbs = []
    for p in _DIR_HISTORY_BACK[-5:]:
        crumbs.append(p.name or str(p))
    crumbs.append(f"[{_CURRENT_DIR.name or str(_CURRENT_DIR)} (current)]")
    for p in reversed(_DIR_HISTORY_FORWARD[-5:]):
        crumbs.append(p.name or str(p))
    return "Navigation path: " + " -> ".join(crumbs)


# ── DRIVES ───────────────────────────────────────────────────────────────────

def list_drives() -> str:
    """List all drives with labels, types, total/free space and usage bars."""
    drives = []
    if _OS == "Windows" and _WIN32_AVAILABLE:
        try:
            for d in win32api.GetLogicalDriveStrings().split('\000')[:-1]:
                try:
                    dtype_code = win32file.GetDriveType(d)
                    type_names = {
                        win32file.DRIVE_REMOVABLE: "Removable USB",
                        win32file.DRIVE_FIXED:     "Fixed Drive (SSD/HDD)",
                        win32file.DRIVE_REMOTE:    "Network Drive",
                        win32file.DRIVE_CDROM:     "CD/DVD Drive",
                        win32file.DRIVE_RAMDISK:   "RAM Disk",
                    }
                    dtype = type_names.get(dtype_code, "Local Drive")
                    try:
                        label = win32api.GetVolumeInformation(d)[0] or "Local Disk"
                    except Exception:
                        label = "Local Disk"

                    usage = shutil.disk_usage(d)
                    used_pct = (usage.used / usage.total * 100) if usage.total > 0 else 0
                    drives.append({
                        "letter": d, "label": label, "type": dtype,
                        "total": usage.total, "used": usage.used,
                        "free": usage.free, "pct": used_pct,
                    })
                except Exception:
                    continue
        except Exception:
            pass

    if not drives:
        check_paths = [Path("/"), Path.home()]
        if _OS == "Windows":
            check_paths += [Path(f"{c}:/") for c in "CDEFGH"]
        seen = set()
        for p in check_paths:
            try:
                if p.exists() and str(p.resolve()) not in seen:
                    seen.add(str(p.resolve()))
                    usage = shutil.disk_usage(p)
                    pct = (usage.used / usage.total * 100) if usage.total > 0 else 0
                    drives.append({
                        "letter": str(p), "label": p.name or "Root", "type": "Storage",
                        "total": usage.total, "used": usage.used,
                        "free": usage.free, "pct": pct,
                    })
            except Exception:
                continue

    if not drives:
        return "Could not retrieve drive information."

    lines = [f"System Storage & Drives ({len(drives)} available):"]
    for d in drives:
        bar_len = 14
        filled = int(bar_len * d["pct"] / 100)
        bar = "=" * filled + "-" * (bar_len - filled)
        status = "OK" if d["pct"] < 85 else ("WARNING" if d["pct"] < 95 else "CRITICAL")
        lines.append(
            f"  {d['letter']} [{d['label']}] ({d['type']})\n"
            f"    [{bar}] {d['pct']:.1f}% used | Free: {_format_size(d['free'])} / {_format_size(d['total'])} ({status})"
        )
    return "\n".join(lines)


# ── TREE VIEW ────────────────────────────────────────────────────────────────

def tree_view(path: str = "", max_depth: int = 2, max_items: int = 40) -> str:
    target = _resolve_path(path) if path else _CURRENT_DIR
    if not target.exists():
        return f"Path not found: {target}"
    if not target.is_dir():
        return f"Not a directory: {target}"
    if not _is_safe_path(target):
        return f"Access denied: {target}"

    lines = [f"{target.name or str(target)}/"]
    item_count = [0]

    def _walk(p: Path, prefix: str, current_depth: int):
        if current_depth > max_depth or item_count[0] >= max_items:
            return
        try:
            entries = sorted(list(p.iterdir()), key=lambda x: (not x.is_dir(), x.name.lower()))
        except Exception:
            return

        visible = [e for e in entries if not e.name.startswith(".")]
        total = len(visible)
        for i, entry in enumerate(visible):
            if item_count[0] >= max_items:
                lines.append(f"{prefix}... [truncated at {max_items} items]")
                break
            item_count[0] += 1
            is_last = (i == total - 1)
            connector = "\\-- " if is_last else "|-- "
            sub_prefix = "    " if is_last else "|   "
            if entry.is_dir():
                lines.append(f"{prefix}{connector}{entry.name}/")
                _walk(entry, prefix + sub_prefix, current_depth + 1)
            else:
                try:
                    sz = _format_size(entry.stat().st_size)
                    lines.append(f"{prefix}{connector}{entry.name} ({sz})")
                except Exception:
                    lines.append(f"{prefix}{connector}{entry.name}")

    _walk(target, "", 1)
    return "\n".join(lines)


# ── IN-FILE CONTENT SEARCH ───────────────────────────────────────────────────

def search_file_content(path: str = "", query: str = "", extension: str = "", max_results: int = 15) -> str:
    if not query:
        return _need("What text should I search for inside the files?")
    target = _resolve_path(path) if path else _CURRENT_DIR
    if not target.exists():
        return f"Search path not found: {target}"
    if not _is_safe_path(target):
        return f"Access denied: {target}"

    results = []
    clean_q = query.lower().strip()
    scanned_count = 0
    max_scan_files = 300
    ext_filter = ("." + extension.lstrip(".")).lower() if extension else ""
    binary_exts = {
        ".pyc", ".pyd", ".exe", ".dll", ".so", ".dylib", ".bin", ".obj",
        ".o", ".class", ".zip", ".tar", ".gz", ".7z", ".rar", ".iso",
        ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".ico",
        ".mp4", ".mkv", ".mov", ".avi", ".mp3", ".wav", ".flac", ".ogg"
    }

    for item in target.rglob("*"):
        if scanned_count >= max_scan_files or len(results) >= max_results:
            break
        if not item.is_file() or item.name.startswith("."):
            continue
        if ext_filter and item.suffix.lower() != ext_filter:
            continue
        if item.suffix.lower() in binary_exts:
            continue
        try:
            if item.stat().st_size > 10 * 1024 * 1024:
                continue
        except Exception:
            continue

        scanned_count += 1
        try:
            with open(item, "r", encoding="utf-8", errors="ignore") as f:
                for line_num, line in enumerate(f, 1):
                    if clean_q in line.lower():
                        snippet = line.strip()
                        if len(snippet) > 130:
                            snippet = snippet[:127] + "..."
                        results.append(f"  {item.name}:{line_num} -> {snippet}  ({item.parent})")
                        if len(results) >= max_results:
                            break
        except Exception:
            continue

    if not results:
        return f"No matches found for '{query}' in {target.name}/ (scanned {scanned_count} files)."

    return f"Found {len(results)} match(es) for '{query}' in {target.name}/:\n" + "\n".join(results)


# ── RECENT FILES ─────────────────────────────────────────────────────────────

def get_recent_files(path: str = "downloads", count: int = 10, hours: float = 0) -> str:
    target = _resolve_path(path)
    if not target.exists():
        return f"Path not found: {target}"
    if not _is_safe_path(target):
        return f"Access denied: {target}"

    count = min(max(count, 1), 50)
    files = []
    cutoff_ts = (time.time() - hours * 3600) if hours > 0 else 0

    try:
        for item in target.iterdir():
            if item.is_file() and not item.name.startswith("."):
                try:
                    mtime = item.stat().st_mtime
                    if mtime >= cutoff_ts:
                        files.append((mtime, item))
                except Exception:
                    continue
    except Exception as e:
        return f"Could not inspect recent files: {e}"

    files.sort(key=lambda x: x[0], reverse=True)
    top = files[:count]

    if not top:
        time_msg = f" in the last {hours} hours" if hours > 0 else ""
        return f"No files found in {target.name}/{time_msg}."

    lines = [f"Recent files in {target.name}/ (newest first):"]
    for mtime, f in top:
        dt_str = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M")
        size_str = _format_size(f.stat().st_size)
        lines.append(f"  {dt_str} | {size_str:>9} | {f.name}")
    return "\n".join(lines)


# ── FOLDER ORGANIZER ─────────────────────────────────────────────────────────

def organize_folder(path: str = "") -> str:
    """Organize a folder into categorized subfolders. Asks which folder if none given."""
    if not str(path or "").strip():
        return _need("Which folder should I organize? (e.g. Desktop, Downloads, D:/stuff)")

    type_map = {
        "Images":     {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".svg", ".ico", ".heic", ".raw"},
        "Documents":  {".pdf", ".doc", ".docx", ".txt", ".rtf", ".odt", ".xls", ".xlsx", ".csv", ".ppt", ".pptx"},
        "Videos":     {".mp4", ".avi", ".mkv", ".mov", ".wmv", ".flv", ".webm", ".m4v"},
        "Music":      {".mp3", ".wav", ".flac", ".aac", ".ogg", ".wma", ".m4a", ".opus"},
        "Archives":   {".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz", ".iso"},
        "Code":       {".py", ".js", ".ts", ".html", ".css", ".json", ".xml", ".cpp", ".c", ".h", ".cs", ".go", ".rs", ".sql", ".sh", ".bat", ".ps1"},
        "Installers": {".exe", ".msi", ".dmg", ".pkg", ".appx"},
    }

    target_dir = _resolve_path(path)
    if not target_dir.exists() or not target_dir.is_dir():
        return _need(f"The folder '{target_dir}' was not found. Which folder do you mean?")
    if not _is_safe_path(target_dir):
        return f"Access denied: {target_dir}"

    moved, skipped = [], []
    journal: list[tuple[Path, Path]] = []

    try:
        category_names = set(type_map.keys()) | {"Others"}
        for item in target_dir.iterdir():
            if item.is_dir() or item.name.startswith("."):
                continue
            if item.name in category_names:
                continue

            ext = item.suffix.lower()
            dest_folder = target_dir / "Others"
            for cat, exts in type_map.items():
                if ext in exts:
                    dest_folder = target_dir / cat
                    break

            dest_folder.mkdir(exist_ok=True)
            new_path = dest_folder / item.name

            if new_path.exists():
                skipped.append(item.name)
                continue

            origin = item.resolve()
            shutil.move(str(item), str(new_path))
            journal.append((origin, new_path.resolve()))
            moved.append(f"{item.name} -> {dest_folder.name}/")

        if journal:
            def _undo_organize(entries=tuple(journal)):
                restored = 0
                for origin, moved_to in entries:
                    try:
                        if moved_to.exists():
                            origin.parent.mkdir(parents=True, exist_ok=True)
                            shutil.move(str(moved_to), str(origin))
                            restored += 1
                    except Exception as e:
                        print(f"[file] undo organize error: {e}")
                for folder in {m.parent for _o, m in entries}:
                    try:
                        if folder.exists() and folder.is_dir() and not any(folder.iterdir()):
                            folder.rmdir()
                    except Exception:
                        pass
                return f"{restored} file(s) restored to {target_dir.name}/."
            push_undo(f"organized {target_dir.name}/ ({len(journal)} files)", _undo_organize)

        result = f"Organized {target_dir.name}/: {len(moved)} files moved."
        if moved:
            preview = moved[:8]
            result += "\n" + "\n".join(f"  {m}" for m in preview)
            if len(moved) > 8:
                result += f"\n  ... and {len(moved) - 8} more."
        if skipped:
            result += f"\n  {len(skipped)} file(s) skipped due to name conflict."
        return result
    except Exception as e:
        return f"Could not organize folder: {e}"


# ── DUPLICATES ───────────────────────────────────────────────────────────────

def find_duplicate_files(path: str = "downloads", max_results: int = 15) -> str:
    target = _resolve_path(path)
    if not target.exists():
        return f"Path not found: {target}"
    if not _is_safe_path(target):
        return f"Access denied: {target}"

    by_size = defaultdict(list)
    scanned = 0
    max_scan = 1000

    for item in target.rglob("*"):
        if scanned >= max_scan:
            break
        if not item.is_file() or item.name.startswith("."):
            continue
        try:
            sz = item.stat().st_size
            if sz > 0:
                by_size[sz].append(item)
                scanned += 1
        except Exception:
            continue

    potential = {sz: paths for sz, paths in by_size.items() if len(paths) > 1}
    duplicates = []
    wasted_bytes = 0

    for sz, paths in potential.items():
        by_hash = defaultdict(list)
        for p in paths:
            try:
                h = hashlib.md5()
                with open(p, "rb") as f:
                    while chunk := f.read(65536):   # hash the WHOLE file
                        h.update(chunk)
                by_hash[h.hexdigest()].append(p)
            except Exception:
                continue

        for h, dups in by_hash.items():
            if len(dups) > 1:
                wasted_bytes += sz * (len(dups) - 1)
                duplicates.append((sz, dups))
                if len(duplicates) >= max_results:
                    break
        if len(duplicates) >= max_results:
            break

    if not duplicates:
        return f"No duplicate files detected in {target.name}/ (scanned {scanned} files)."

    lines = [f"Found {len(duplicates)} duplicate group(s) in {target.name}/ (Recoverable: {_format_size(wasted_bytes)}):"]
    for i, (sz, dups) in enumerate(duplicates, 1):
        lines.append(f"  Group #{i} ({_format_size(sz)} each):")
        for p in dups:
            lines.append(f"    - {p.name}  ({p.parent})")
    return "\n".join(lines)


# ── COMPRESSION ──────────────────────────────────────────────────────────────

def compress_target(path: str, name: str = "", destination: str = "", format: str = "zip") -> str:
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("What file or folder should I compress, and where is it?")
        base = _resolve_path(path)
        src = (base / name) if name else base
        if not src.exists():
            return _need(f"I could not find '{src}'. Which file/folder do you mean?")
        if not _is_safe_path(src):
            return f"Access denied: {src}"

        if destination:
            dst = _resolve_path(destination)
            if dst.is_dir():
                dst = dst / f"{src.stem}.zip"
        else:
            dst = src.parent / f"{src.stem}.zip"

        dst.parent.mkdir(parents=True, exist_ok=True)

        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
            if src.is_dir():
                for root, _, files in os.walk(src):
                    for file in files:
                        full_p = Path(root) / file
                        if full_p == dst:
                            continue
                        z.write(full_p, arcname=full_p.relative_to(src.parent))
            else:
                z.write(src, arcname=src.name)

        if not dst.is_file():
            return f"FAILED: archive was not created at {dst}"

        def _undo_zip():
            if dst.exists():
                dst.unlink()
                return f"Removed created archive '{dst.name}'."
            return f"Archive '{dst.name}' already gone."
        push_undo(f"compressed {src.name} to {dst.name}", _undo_zip)

        return f"VERIFIED: compressed to {dst.resolve()} ({_format_size(dst.stat().st_size)})"
    except Exception as e:
        return f"Compression failed: {e}"


def _safe_extract_zip(z: zipfile.ZipFile, dst: Path):
    dst_res = dst.resolve()
    for m in z.namelist():
        if not (dst_res / m).resolve().is_relative_to(dst_res):
            raise ValueError(f"Unsafe path in archive: {m}")
    z.extractall(dst)


def _safe_extract_tar(t: tarfile.TarFile, dst: Path):
    dst_res = dst.resolve()
    for m in t.getmembers():
        if not (dst_res / m.name).resolve().is_relative_to(dst_res):
            raise ValueError(f"Unsafe path in archive: {m.name}")
    t.extractall(dst)


def extract_archive(path: str, name: str = "", destination: str = "") -> str:
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("Which archive should I extract, and where is it?")
        base = _resolve_path(path)
        src = (base / name) if name else base
        if not src.exists():
            return _need(f"I could not find the archive '{src}'. Where is it?")
        if not _is_safe_path(src):
            return f"Access denied: {src}"

        dst = _resolve_path(destination) if destination else src.parent / src.stem
        if not _is_safe_path(dst):
            return f"Access denied: {dst}"
        dst.mkdir(parents=True, exist_ok=True)

        if zipfile.is_zipfile(src):
            with zipfile.ZipFile(src, "r") as z:
                _safe_extract_zip(z, dst)
        elif tarfile.is_tarfile(src):
            with tarfile.open(src, "r:*") as t:
                _safe_extract_tar(t, dst)
        else:
            return f"Unsupported archive format: {src.suffix}"

        return f"Extracted '{src.name}' into: {dst.resolve()}"
    except Exception as e:
        return f"Extraction failed: {e}"


# ── HEAD / TAIL / CHECKSUM ───────────────────────────────────────────────────

def file_head_tail(path: str, name: str = "", lines: int = 25, from_end: bool = False) -> str:
    try:
        base = _resolve_path(path)
        target = (base / name) if name else base
        if not target.exists() or not target.is_file():
            return _need(f"File not found: '{target}'. Which file and folder do you mean?")
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        with open(target, "r", encoding="utf-8", errors="ignore") as f:
            all_lines = f.readlines()

        selected = all_lines[-lines:] if from_end else all_lines[:lines]
        label = "Last" if from_end else "First"
        header = f"{label} {len(selected)} lines of {target.name} (Total: {len(all_lines)} lines):\n"
        return header + "".join(selected)
    except Exception as e:
        return f"Could not read lines from file: {e}"


def get_file_checksum(path: str, name: str = "", algorithm: str = "sha256") -> str:
    try:
        base = _resolve_path(path)
        target = (base / name) if name else base
        if not target.exists() or not target.is_file():
            return _need(f"File not found: '{target}'. Which file and folder do you mean?")
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        h = hashlib.sha256() if algorithm.lower() == "sha256" else hashlib.md5()
        with open(target, "rb") as f:
            while chunk := f.read(65536):
                h.update(chunk)

        return f"{algorithm.upper()} checksum for '{target.name}':\n  {h.hexdigest()}"
    except Exception as e:
        return f"Could not calculate checksum: {e}"


# ── RECYCLE BIN ──────────────────────────────────────────────────────────────

def get_recycle_bin_info() -> str:
    if _OS != "Windows":
        return "Recycle Bin query is only supported on Windows."
    try:
        import ctypes
        from ctypes import wintypes

        class SHQUERYRBINFO(ctypes.Structure):
            _fields_ = [
                ("cbSize", wintypes.DWORD),
                ("i64Size", ctypes.c_int64),
                ("i64NumItems", ctypes.c_int64),
            ]

        info = SHQUERYRBINFO()
        info.cbSize = ctypes.sizeof(SHQUERYRBINFO)
        res = ctypes.windll.shell32.SHQueryRecycleBinW(None, ctypes.byref(info))
        if res == 0:
            return f"Recycle Bin status:\n  Items : {info.i64NumItems}\n  Total Size : {_format_size(info.i64Size)}"
        return "Could not query Recycle Bin information."
    except Exception as e:
        return f"Recycle Bin error: {e}"


def empty_recycle_bin(confirmed: bool = False) -> str:
    """Permanently empties the Recycle Bin. Requires explicit confirmation."""
    if _OS != "Windows":
        return "Emptying trash is only implemented on Windows."
    if not confirmed:
        return _need("Emptying the Recycle Bin is PERMANENT. Do you really want to empty it? (yes/no)")
    try:
        import ctypes
        res = ctypes.windll.shell32.SHEmptyRecycleBinW(None, None, 0x00000007)
        if res == 0:
            return "Recycle Bin has been completely emptied."
        return "Recycle Bin is already empty or could not be cleared."
    except Exception as e:
        return f"Could not empty Recycle Bin: {e}"


# ── STANDARD FILE OPERATIONS ─────────────────────────────────────────────────

_TEXT_EXTS = {".txt", ".md", ".py", ".js", ".ts", ".html", ".css", ".json", ".xml",
              ".csv", ".log", ".ini", ".cfg", ".yaml", ".yml", ".toml", ".bat", ".ps1",
              ".sh", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".sql", ".java", ".rtf",
              ".pdf", ".docx", ".xlsx", ".xls", ".ipynb"}


def open_file(path: str = "", name: str = "", read_content: bool = True) -> str:
    """Open a file with its default app. If no folder is given, searches Desktop,
    Downloads, Documents and all other drives, and asks when unsure."""
    try:
        global _CURRENT_DIR
        path, name = str(path or "").strip(), str(name or "").strip()
        if not name and not path:
            return _need("Which file do you want to open, and in which folder?")

        if path:
            base = _resolve_path(path)
            if name:
                if not base.is_dir():
                    return _need(f"'{base}' is not a folder. Which folder is '{name}' in?")
                target = base / name
                if not target.exists():
                    # try a partial-name match inside that folder only
                    hits = [p for p in base.iterdir() if name.lower() in p.name.lower()]
                    if len(hits) == 1:
                        target = hits[0]
                    elif len(hits) > 1:
                        listing = "\n".join(f"  {i+1}. {m.name}" for i, m in enumerate(hits[:10]))
                        return _need(f"Several items in '{base}' match '{name}':\n{listing}\nWhich one?")
                    else:
                        return _need(f"'{name}' is not in '{base}'. Check the name, or should I search all drives?")
            else:
                target = base
                if not target.exists():
                    return _need(f"I could not find '{target}'. Check the path, or should I search all drives?")
        else:
            matches = _search_everywhere(name)
            if not matches:
                return _need(f"No file named '{name}' found on Desktop, Downloads, Documents or other drives. Which folder is it in?")
            if len(matches) > 1:
                listing = "\n".join(f"  {i+1}. {m}" for i, m in enumerate(matches))
                return _need(f"I found {len(matches)} matches:\n{listing}\nWhich one should I open?")
            target = matches[0]

        if target.is_dir():
            return open_folder(str(target))
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        if _OS == "Windows":
            os.startfile(str(target))
        elif _OS == "Darwin":
            subprocess.Popen(["open", str(target)])
        else:
            subprocess.Popen(["xdg-open", str(target)])

        _CURRENT_DIR = target.parent
        msg = f"Opened: {target.resolve()}"
        if read_content and target.suffix.lower() in _TEXT_EXTS:
            msg += f"\n\nContent:\n{read_file(str(target))}"
        return msg
    except Exception as e:
        return f"Error opening file: {e}"


def list_files(path: str = "desktop", show_hidden: bool = False) -> str:
    try:
        target = _resolve_path(path)
        if not _is_safe_path(target):
            return f"Access denied: {target}"
        if not target.exists():
            return _need(f"The folder '{target}' does not exist. Which folder do you mean?")
        if not target.is_dir():
            return f"Not a directory: {target}"

        items = []
        for item in sorted(target.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower())):
            if not show_hidden and item.name.startswith("."):
                continue
            if item.is_dir():
                items.append(f"📁 {item.name}/")
            else:
                try:
                    size = _format_size(item.stat().st_size)
                except Exception:
                    size = "?"
                items.append(f"📄 {item.name} ({size})")

        if not items:
            return f"Directory is empty: {target}"

        return f"Contents of {target} ({len(items)} items):\n" + "\n".join(items)
    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Error listing files: {e}"


def create_file(path: str = "", name: str = "", content=None, overwrite: bool = False) -> str:
    """Create a file. NEVER guesses: asks for location, name and content first,
    and only reports success after reading the file back from disk."""
    try:
        global _CURRENT_DIR
        path = str(path or "").strip().strip('"').strip("'")
        name = str(name or "").strip().strip('"').strip("'")

        if not path:
            return _need("Where should I create the file? (e.g. Desktop, Downloads, D:/projects)")
        if not name:
            return _need("What should the file be called (with extension, e.g. notes.txt)?")
        if content is None:
            return _need(f"What should be written inside '{name}'? (say 'empty' for a blank file)")
        content = str(content)
        if content.strip().lower() == "empty":
            content = ""

        if any(c in name for c in '<>:"|?*') or name in (".", ".."):
            return _need(f"'{name}' is not a valid file name. What name should I use?")

        folder = _resolve_path(path)
        if folder.is_file():
            return _need(f"'{path}' is a file, not a folder. Which folder should I use?")
        if not folder.exists():
            return _need(f"The folder '{folder}' does not exist. Should I create it, or use another location?")

        # allow name like "sub/notes.txt"
        target = folder / name
        if not _is_safe_path(target):
            return f"Access denied: {target}"
        if target.exists() and not overwrite:
            return _need(f"'{target}' already exists. Overwrite it, or use a different name?")

        previous = None
        existed = target.exists()
        if existed:
            try:
                previous = target.read_text(encoding="utf-8", errors="ignore")
            except Exception:
                previous = None

        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")

        # VERIFY - never claim success without checking
        if not target.is_file():
            return f"FAILED: file was NOT created at {target}"
        if target.read_text(encoding="utf-8", errors="ignore") != content:
            return f"FAILED: file exists at {target} but content does not match."

        _CURRENT_DIR = target.parent
        push_undo(
            f"created {target.name}",
            _undo_write(target, previous) if existed else _undo_create(target),
        )
        return f"VERIFIED: created {target.resolve()} ({len(content)} characters)."
    except Exception as e:
        return f"FAILED to create file: {e}"


def create_folder(path: str = "", name: str = "") -> str:
    try:
        if not str(path or "").strip():
            return _need("Where should I create the folder? (e.g. Desktop, D:/projects)")
        if not str(name or "").strip():
            return _need("What should the new folder be called?")
        base = _resolve_path(path)
        target = base / name
        if not _is_safe_path(target):
            return f"Access denied: {target}"
        if target.exists():
            return f"Folder already exists: {target.resolve()}"
        target.mkdir(parents=True, exist_ok=True)
        if not target.is_dir():
            return f"FAILED: folder was NOT created at {target}"
        push_undo(f"created folder {target.name}", _undo_create(target))
        return f"VERIFIED: folder created at {target.resolve()}"
    except Exception as e:
        return f"FAILED to create folder: {e}"


def delete_file(path: str = "", name: str = "") -> str:
    """Move a file or folder to the Recycle Bin (with undo)."""
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("What should I delete, and in which folder?")
        base = _resolve_path(path)
        target = (base / name) if name else base
        if not _is_safe_path(target):
            return f"Access denied: {target}"
        if not target.exists():
            return _need(f"I could not find '{target}'. Which file/folder do you mean?")

        protected = {
            _get_desktop(), _get_downloads(), _get_documents(),
            _get_pictures(), _get_music(), _get_videos(), Path.home(),
            Path("C:/"), Path("D:/")
        }
        if target.resolve() in {p.resolve() for p in protected}:
            return f"Protected directory, cannot delete: {target.name}"

        original = target.resolve()
        result = _safe_trash(target)
        if result.startswith("Moved to Trash"):
            if original.exists():
                return f"FAILED: '{original}' is still there."
            push_undo(f"deleted {original.name}", lambda p=original: _restore_from_trash(p))
        return result
    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Could not delete: {e}"


def move_file(path: str = "", name: str = "", destination: str = "") -> str:
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("What should I move, and where is it?")
        if not str(destination or "").strip():
            return _need("Where should I move it to?")
        base = _resolve_path(path)
        src = (base / name) if name else base
        dst = _resolve_path(destination)

        if not src.exists():
            return _need(f"I could not find '{src}'. Which file/folder do you mean?")
        if not _is_safe_path(src):
            return f"Access denied (source): {src}"
        if not _is_safe_path(dst):
            return f"Access denied (destination): {dst}"

        if dst.is_dir():
            dst = dst / src.name
        if dst.exists():
            return _need(f"'{dst}' already exists. Use a different destination or name?")

        dst.parent.mkdir(parents=True, exist_ok=True)
        origin = src.resolve()
        shutil.move(str(src), str(dst))
        if not dst.exists():
            return f"FAILED: '{dst}' does not exist after the move."
        push_undo(f"moved {origin.name} to {dst.parent.name}/", _undo_move(origin, dst.resolve()))
        return f"VERIFIED: moved {origin} -> {dst.resolve()}"
    except Exception as e:
        return f"Could not move: {e}"


def copy_file(path: str = "", name: str = "", destination: str = "") -> str:
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("What should I copy, and where is it?")
        if not str(destination or "").strip():
            return _need("Where should I copy it to?")
        base = _resolve_path(path)
        src = (base / name) if name else base
        dst = _resolve_path(destination)

        if not src.exists():
            return _need(f"I could not find '{src}'. Which file/folder do you mean?")
        if not _is_safe_path(src):
            return f"Access denied (source): {src}"
        if not _is_safe_path(dst):
            return f"Access denied (destination): {dst}"

        if dst.is_dir():
            dst = dst / src.name
        if dst.exists():
            return _need(f"'{dst}' already exists. Use a different destination or name?")

        dst.parent.mkdir(parents=True, exist_ok=True)
        if src.is_dir():
            shutil.copytree(str(src), str(dst))
        else:
            shutil.copy2(str(src), str(dst))

        if not dst.exists():
            return f"FAILED: copy was not created at {dst}"

        _copy = dst.resolve()

        def _undo_copy():
            if not _copy.exists():
                return f"The copy '{_copy.name}' is already gone."
            if _copy.is_dir():
                shutil.rmtree(_copy)
            else:
                _copy.unlink()
            return f"Removed the copy in {_copy.parent.name}/."
        push_undo(f"copied {src.name} to {dst.parent.name}/", _undo_copy)
        return f"VERIFIED: copied {src.resolve()} -> {_copy}"
    except Exception as e:
        return f"Could not copy: {e}"


def rename_file(path: str = "", name: str = "", new_name: str = "") -> str:
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("What should I rename, and where is it?")
        if not str(new_name or "").strip():
            return _need("What should the new name be?")
        base = _resolve_path(path)
        target = (base / name) if name else base
        if not _is_safe_path(target):
            return f"Access denied: {target}"
        if not target.exists():
            return _need(f"I could not find '{target}'. Which file/folder do you mean?")

        new_path = target.parent / new_name
        if new_path.exists():
            return _need(f"A file named '{new_name}' already exists there. Use a different name?")

        old_path = target.resolve()
        target.rename(new_path)
        if not new_path.exists():
            return f"FAILED: '{new_path}' does not exist after rename."
        push_undo(f"renamed {old_path.name} to {new_name}", _undo_move(old_path, new_path.resolve()))
        return f"VERIFIED: renamed {old_path.name} -> {new_name} (in {new_path.parent})"
    except Exception as e:
        return f"Could not rename: {e}"


def read_file(path: str = "", name: str = "", max_chars: int = 5000, open_viewer: bool = False) -> str:
    """Read text from PDF, Word, Excel, Jupyter, text and code files.
    If no folder is given, searches all drives for the file name."""
    try:
        path, name = str(path or "").strip(), str(name or "").strip()
        if not path and not name:
            return _need("Which file should I read, and in which folder?")

        if path:
            base = _resolve_path(path)
            target = (base / name) if name else base
            if target.is_dir() and not name:
                return _need(f"'{target}' is a folder. Which file inside it should I read?")
            if not target.exists() and name and base.is_dir():
                hits = [p for p in base.iterdir() if p.is_file() and name.lower() in p.name.lower()]
                if len(hits) == 1:
                    target = hits[0]
                elif len(hits) > 1:
                    listing = "\n".join(f"  {i+1}. {m.name}" for i, m in enumerate(hits[:10]))
                    return _need(f"Several files in '{base}' match '{name}':\n{listing}\nWhich one?")
        else:
            matches = _search_everywhere(name)
            if not matches:
                return _need(f"No file named '{name}' found. Which folder is it in?")
            if len(matches) > 1:
                listing = "\n".join(f"  {i+1}. {m}" for i, m in enumerate(matches))
                return _need(f"I found {len(matches)} matches:\n{listing}\nWhich one should I read?")
            target = matches[0]

        if not target.exists():
            return _need(f"File not found: '{target}'. Check the name and folder.")
        if not target.is_file():
            return f"Not a file: {target.name}"
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        if open_viewer and _OS == "Windows":
            try:
                os.startfile(str(target))
            except Exception:
                pass

        ext = target.suffix.lower()

        # PDF
        if ext == ".pdf":
            extracted = ""
            try:
                import pdfplumber
                with pdfplumber.open(target) as pdf:
                    for i, page in enumerate(pdf.pages):
                        txt = page.extract_text() or ""
                        if txt.strip():
                            extracted += f"--- Page {i+1} ---\n{txt}\n\n"
                        if len(extracted) >= max_chars:
                            break
            except Exception:
                try:
                    import PyPDF2
                    with open(target, "rb") as f:
                        reader = PyPDF2.PdfReader(f)
                        for i, page in enumerate(reader.pages):
                            txt = page.extract_text() or ""
                            if txt.strip():
                                extracted += f"--- Page {i+1} ---\n{txt}\n\n"
                            if len(extracted) >= max_chars:
                                break
                except Exception as e:
                    return f"Error reading PDF: {e}"

            if not extracted.strip():
                return f"PDF '{target.name}' contains no extractable text (may be scanned)."
            if len(extracted) > max_chars:
                extracted = extracted[:max_chars] + f"\n\n[Truncated — {len(extracted)} total chars]"
            return f"Read from {target}:\n\n{extracted}"

        # Word
        if ext == ".docx":
            try:
                import docx
                doc = docx.Document(target)
                doc_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                if not doc_text.strip():
                    return f"Word document '{target.name}' is empty."
                if len(doc_text) > max_chars:
                    doc_text = doc_text[:max_chars] + "\n\n[Truncated]"
                return f"Read from {target}:\n\n{doc_text}"
            except Exception as e:
                return f"Error reading Word document: {e}"

        # Excel (.xlsx via openpyxl; old .xls needs xlrd)
        if ext in (".xlsx", ".xls"):
            try:
                if ext == ".xls":
                    import xlrd
                    wb = xlrd.open_workbook(str(target))
                    sh = wb.sheet_by_index(0)
                    rows = []
                    for r in range(min(sh.nrows, 20)):
                        vals = [str(c) for c in sh.row_values(r) if str(c).strip()]
                        if vals:
                            rows.append(" | ".join(vals))
                    return f"Read from {target}:\n\nSheets: {', '.join(wb.sheet_names())}\n" + "\n".join(rows[:15])
                import openpyxl
                wb = openpyxl.load_workbook(target, read_only=True, data_only=True)
                summary = f"Spreadsheet with sheets: {', '.join(wb.sheetnames)}\n"
                sheet = wb.active
                rows = []
                for r in sheet.iter_rows(max_row=20, values_only=True):
                    if any(c is not None for c in r):
                        rows.append(" | ".join(str(c) for c in r if c is not None))
                summary += "\n".join(rows[:15])
                return f"Read from {target}:\n\n{summary}"
            except Exception as e:
                return f"Error reading spreadsheet: {e}"

        # Jupyter
        if ext == ".ipynb":
            try:
                nb = json.loads(target.read_text(encoding="utf-8", errors="ignore"))
                cells_text = []
                for idx, cell in enumerate(nb.get("cells", []), 1):
                    ctype = cell.get("cell_type", "code")
                    src = "".join(cell.get("source", []))
                    if src.strip():
                        cells_text.append(f"[{ctype.upper()} CELL {idx}]:\n{src}")
                    if sum(len(c) for c in cells_text) >= max_chars:
                        break
                joined = "\n\n".join(cells_text)
                if len(joined) > max_chars:
                    joined = joined[:max_chars] + "\n\n[Truncated]"
                return f"Read from Jupyter Notebook {target}:\n\n{joined}"
            except Exception as e:
                return f"Error parsing Jupyter notebook: {e}"

        # Text / code
        try:
            content = target.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            try:
                content = target.read_text(encoding="latin-1", errors="ignore")
            except Exception as e:
                return f"Could not read file: {e}"

        if len(content) > max_chars:
            content = content[:max_chars] + f"\n\n[Truncated — {len(content)} total chars]"
        return content
    except Exception as e:
        return f"Could not read file: {e}"


def write_file(path: str = "", name: str = "", content=None, append: bool = False) -> str:
    """Write or append text to an EXISTING-or-new file. Asks for anything missing; verifies result."""
    try:
        if not str(path or "").strip():
            return _need("Which folder is the file in (or should it go in)?")
        if not str(name or "").strip():
            return _need("Which file should I write to?")
        if content is None:
            return _need(f"What text should I {'append to' if append else 'write into'} '{name}'?")
        content = str(content)

        base = _resolve_path(path)
        if not base.is_dir():
            return _need(f"The folder '{base}' does not exist. Which folder do you mean?")
        target = base / name
        if not _is_safe_path(target):
            return f"Access denied: {target}"

        previous: str | None = None
        undoable = True
        if target.exists():
            try:
                if target.stat().st_size > _UNDO_CONTENT_LIMIT:
                    undoable = False
                else:
                    previous = target.read_text(encoding="utf-8", errors="ignore")
            except Exception:
                undoable = False

        mode = "a" if append else "w"
        with open(target, mode, encoding="utf-8") as f:
            f.write(content)

        if not target.is_file():
            return f"FAILED: '{target}' does not exist after writing."
        if not append and target.read_text(encoding="utf-8", errors="ignore") != content:
            return f"FAILED: content in '{target}' does not match what was requested."

        action = "appended to" if append else "written to"
        if undoable:
            push_undo(f"wrote to {target.name}", _undo_write(target, previous))
            return f"VERIFIED: {action} {target.resolve()}"
        return f"VERIFIED: {action} {target.resolve()} (too large for in-memory undo)"
    except Exception as e:
        return f"Could not write file: {e}"


def find_files(name: str = "", extension: str = "", path: str = "", max_results: int = 20) -> str:
    """Find files by name/extension. With no folder given, searches Desktop, Downloads,
    Documents and every other drive."""
    try:
        if not name and not extension:
            return _need("What file name or extension should I look for?")

        if path:
            search_path = _resolve_path(path)
            if not _is_safe_path(search_path):
                return f"Access denied: {search_path}"
            if not search_path.exists():
                return _need(f"The folder '{search_path}' does not exist. Which folder should I search?")
            roots = [search_path]
            where = str(search_path)
        else:
            roots = None
            where = "Desktop, Downloads, Documents and other drives"

        matches = _search_everywhere(name=name, extension=extension, limit=max_results,
                                     time_cap=20.0, roots=roots)
        if not matches:
            return f"No match for '{name or extension}' in {where}."

        lines = []
        for p in matches:
            try:
                size = _format_size(p.stat().st_size)
            except Exception:
                size = "?"
            lines.append(f"📄 {p.name} ({size}) — {p.parent}")
        return f"Found {len(lines)} file(s) in {where}:\n" + "\n".join(lines)
    except Exception as e:
        return f"Search error: {e}"


def get_largest_files(path: str = "downloads", count: int = 10) -> str:
    count = min(max(count, 1), 50)
    try:
        search_path = _resolve_path(path)
        if not _is_safe_path(search_path):
            return f"Access denied: {search_path}"
        if not search_path.exists():
            return f"Path not found: {path}"

        files = []
        for item in search_path.rglob("*"):
            if item.is_file():
                try:
                    files.append((item.stat().st_size, item))
                except Exception:
                    continue

        files.sort(key=lambda x: x[0], reverse=True)
        top = files[:count]

        if not top:
            return "No files found."

        lines = [f"Top {len(top)} largest files in {search_path.name}/:"]
        for size, f in top:
            lines.append(f"  {_format_size(size):>10}  {f.name}  ({f.parent})")
        return "\n".join(lines)
    except Exception as e:
        return f"Error: {e}"


def get_disk_usage(path: str = "home") -> str:
    try:
        target = _resolve_path(path)
        usage = shutil.disk_usage(target)
        pct = usage.used / usage.total * 100
        return (
            f"Disk usage ({target}):\n"
            f"  Total : {_format_size(usage.total)}\n"
            f"  Used  : {_format_size(usage.used)} ({pct:.1f}%)\n"
            f"  Free  : {_format_size(usage.free)}"
        )
    except Exception as e:
        return f"Could not get disk usage: {e}"


def get_file_info(path: str = "", name: str = "") -> str:
    try:
        if not str(path or "").strip() and not str(name or "").strip():
            return _need("Which file or folder do you want info about?")
        base = _resolve_path(path)
        target = (base / name) if name else base
        if not _is_safe_path(target):
            return f"Access denied: {target}"
        if not target.exists():
            return _need(f"I could not find '{target}'. Which file/folder do you mean?")

        stat = target.stat()
        info = {
            "Name":      target.name,
            "Type":      "Folder" if target.is_dir() else "File",
            "Size":      _format_size(stat.st_size),
            "Location":  str(target.parent),
            "Created":   datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M"),
            "Modified":  datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M"),
            "Extension": target.suffix or "—",
        }
        return "\n".join(f"  {k}: {v}" for k, v in info.items())
    except Exception as e:
        return f"Could not get file info: {e}"


# ── MAIN DISPATCHER ──────────────────────────────────────────────────────────

def file_controller(
    parameters: dict = None,
    response=None,
    player=None,
    session_memory=None,
) -> str:
    params = parameters or {}
    action = str(params.get("action", "")).lower().strip()

    raw_path = (
        params.get("path")
        or params.get("file_path")
        or params.get("filepath")
        or params.get("target")
        or params.get("folder")
        or params.get("directory")
        or params.get("location")
        or ""
    )
    raw_name = (
        params.get("name")
        or params.get("filename")
        or params.get("file_name")
        or ""
    )

    path = str(raw_path or "")
    name = str(raw_name or "")

    if player:
        try:
            player.write_log(f"[file] {action} {name or path}")
        except Exception:
            pass

    try:
        # 1. Navigation & Explorer
        if action in ("open_folder", "navigate", "go_to", "open_dir", "cd", "explorer", "open_explorer",
                      "open_new", "open_new_folder", "new_window", "new_explorer"):
            new_win = bool(params.get("new_window", False)) or action in (
                "open_new", "open_new_folder", "new_window", "new_explorer")
            if not path and not new_win:
                return _need("Which folder should I open? (e.g. Desktop, Downloads, D:/projects)")
            return open_folder(path, new_window=new_win)

        elif action in ("reveal", "reveal_in_explorer", "show_in_explorer", "show_in_folder", "select_in_explorer", "locate"):
            return reveal_in_explorer(path=path, name=name, new_window=bool(params.get("new_window", False)))

        elif action in ("back", "navigate_back", "go_back"):
            return navigate_back()

        elif action in ("forward", "navigate_forward", "go_forward"):
            return navigate_forward()

        elif action in ("up", "navigate_up", "parent"):
            return navigate_up()

        elif action in ("history", "nav_history", "breadcrumbs"):
            return get_navigation_history()

        elif action in ("open_explorers", "open_windows", "list_explorers"):
            return get_open_explorers()

        elif action in ("close_explorer", "close_explorers"):
            return close_explorer(path=path)

        # 2. Drives & Storage
        elif action in ("drives", "list_drives", "all_drives", "get_drives", "storage"):
            return list_drives()

        elif action in ("disk_usage", "usage", "storage_usage"):
            return get_disk_usage(path or str(_CURRENT_DIR))

        # 3. Viewing & Inspecting
        elif action in ("open_file", "open", "launch", "view"):
            if path and not name:
                p_res = _resolve_path(path)
                if p_res.is_dir():
                    return open_folder(str(p_res))
            return open_file(path=path, name=name)

        elif action in ("list", "ls", "dir"):
            if not path:
                return _need("Which folder should I list? (e.g. Desktop, Downloads, D:/)")
            return list_files(path)

        elif action in ("tree", "tree_view", "hierarchy", "folder_tree"):
            if not path:
                return _need("Which folder should I show the tree of?")
            return tree_view(
                path=path,
                max_depth=int(params.get("max_depth") or params.get("depth") or 2),
                max_items=int(params.get("max_items", 40))
            )

        elif action in ("read", "read_file"):
            return read_file(path=path, name=name)

        elif action in ("head", "first_lines"):
            return file_head_tail(
                path=path, name=name,
                lines=int(params.get("lines") or params.get("count") or 25),
                from_end=False
            )

        elif action in ("tail", "last_lines"):
            return file_head_tail(
                path=path, name=name,
                lines=int(params.get("lines") or params.get("count") or 25),
                from_end=True
            )

        elif action in ("info", "properties", "stat"):
            return get_file_info(path, name=name)

        elif action in ("checksum", "hash", "sha256", "md5"):
            return get_file_checksum(
                path=path, name=name,
                algorithm=params.get("algorithm", "sha256")
            )

        # 4. Search & Discovery
        elif action in ("find", "find_files", "search"):
            return find_files(
                name=name or str(params.get("query", "")),
                extension=params.get("extension", ""),
                path=path,
                max_results=min(int(params.get("max_results", 20)), 50),
            )

        elif action in ("search_content", "grep", "find_in_files", "search_in_files", "contains"):
            return search_file_content(
                path=path or str(_CURRENT_DIR),
                query=params.get("query") or params.get("content") or name,
                extension=params.get("extension", ""),
                max_results=min(int(params.get("max_results", 15)), 40),
            )

        elif action in ("recent", "recent_files", "latest", "latest_files"):
            return get_recent_files(
                path=path or "downloads",
                count=int(params.get("count", 10)),
                hours=float(params.get("hours", 0)),
            )

        elif action in ("largest", "big_files"):
            return get_largest_files(
                path=path or "downloads",
                count=int(params.get("count", 10)),
            )

        elif action in ("duplicates", "find_duplicates", "dup_files"):
            return find_duplicate_files(
                path=path or "downloads",
                max_results=int(params.get("max_results", 15))
            )

        # 5. Create / modify
        elif action in ("create_file", "create", "make_file", "new_file"):
            return create_file(
                path=path, name=name,
                content=params.get("content"),          # None if not given -> asks the user
                overwrite=bool(params.get("overwrite", False)),
            )

        elif action in ("create_folder", "make_folder", "mkdir"):
            return create_folder(path, name=name)

        elif action == "delete":
            return delete_file(path, name=name)

        elif action == "move":
            return move_file(path, name=name, destination=params.get("destination", ""))

        elif action == "copy":
            return copy_file(path, name=name, destination=params.get("destination", ""))

        elif action == "rename":
            return rename_file(path, name=name, new_name=params.get("new_name", ""))

        elif action in ("write", "write_file", "append"):
            return write_file(
                path, name=name,
                content=params.get("content"),
                append=bool(params.get("append", False)) or action == "append"
            )

        # 6. Organization & Compression
        elif action in ("organize", "organize_folder", "organize_desktop", "clean_folder", "clean_desktop"):
            if action == "organize_desktop" or action == "clean_desktop":
                path = path or "desktop"
            return organize_folder(path=path)

        elif action in ("zip", "compress", "archive"):
            return compress_target(
                path=path, name=name,
                destination=params.get("destination", ""),
                format=params.get("format", "zip")
            )

        elif action in ("unzip", "extract", "decompress"):
            return extract_archive(
                path=path, name=name,
                destination=params.get("destination", "")
            )

        # 7. Recycle Bin
        elif action in ("recycle_bin", "trash_info", "bin_status"):
            return get_recycle_bin_info()

        elif action in ("empty_trash", "empty_recycle_bin"):
            return empty_recycle_bin(confirmed=bool(params.get("confirmed", False)))

        else:
            return f"Unknown action: '{action}'"

    except Exception as e:
        return f"File controller error ({action}): {e}"


# ── TOOL DECLARATION (AUTO-DISCOVERED BY CORE/ACTION_LOADER.PY) ──────────────
TOOL = {
    "name": "file_controller",
    "description": (
        "Universal file system controller & explorer with full Windows integration: "
        "open_folder (opens drive D:, Desktop, or any folder in File Explorer and shows contents), "
        "reveal (opens File Explorer with that file/folder selected), "
        "list_drives, tree, list, recent_files, "
        "search_content / grep (searches text INSIDE files), "
        "find (searches file names on Desktop, Downloads, Documents and ALL drives), "
        "organize_folder, open_file, read (PDF, Word, Excel, Jupyter, code, text), "
        "create_file, create_folder, delete, move, copy, rename, write, zip, unzip, "
        "duplicates, head, tail, checksum, back, up, recycle_bin. "
        "ALWAYS use this tool for file/folder navigation, creation, and explorer control. "
        "RULES: "
        "(0) File Explorer: every open_folder/reveal/back/up re-uses the SAME Explorer window, so opening D: and then a folder inside it just navigates that window. ONLY set new_window=true (or use action open_new_folder) when the user says 'new window', 'open new' or 'another window'. "
        "(1) For create_file you MUST have path (folder), name and content from the user. "
        "If ANY is missing, ask the user - never guess and never default to Desktop. "
        "(2) For open_file/read: if the user gave a folder pass it as path; otherwise leave path "
        "empty so all drives are searched. "
        "(3) If a result starts with NEED_INFO, ask the user that exact question and wait for the answer, then call again. "
        "(4) Only tell the user something was created/moved/copied/renamed/written if the result "
        "starts with VERIFIED. If it starts with FAILED, tell the user it failed and why. "
        "(5) Never pass confirmed=true for empty_trash unless the user explicitly said yes."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "action": {
                "type": "STRING",
                "description": (
                    "open_folder | reveal | list_drives | tree | list | open_file | read | "
                    "create_file | create_folder | delete | move | copy | rename | write | "
                    "find | search_content | recent_files | largest | duplicates | organize_folder | "
                    "zip | unzip | head | tail | checksum | back | forward | up | history | "
                    "recycle_bin | empty_trash | open_explorers | close_explorer | open_new_folder | info"
                )
            },
            "path": {
                "type": "STRING",
                "description": "Folder (or file) path or shortcut the USER named: 'desktop', 'downloads', 'documents', 'D:/', 'D:/folder', 'drive e'. Leave empty if the user did not say - then ask them."
            },
            "name": {
                "type": "STRING",
                "description": "File or folder name (e.g. notes.txt)"
            },
            "destination": {
                "type": "STRING",
                "description": "Destination folder or file path for move, copy, zip, or unzip"
            },
            "new_name": {
                "type": "STRING",
                "description": "New filename for rename action"
            },
            "content": {
                "type": "STRING",
                "description": "Text to put inside the file for create_file/write. Use the word 'empty' for a blank file. Omit if the user has not said."
            },
            "overwrite": {
                "type": "BOOLEAN",
                "description": "Only true if the user agreed to overwrite an existing file"
            },
            "append": {
                "type": "BOOLEAN",
                "description": "For write: add to the end of the file instead of replacing it"
            },
            "new_window": {
                "type": "BOOLEAN",
                "description": "true ONLY if the user asks for a new/another File Explorer window. Default false = re-use the same window"
            },
            "confirmed": {
                "type": "BOOLEAN",
                "description": "For empty_trash: true only after the user explicitly confirmed"
            },
            "query": {
                "type": "STRING",
                "description": "Search keyword or text for search_content / grep"
            },
            "extension": {
                "type": "STRING",
                "description": "File extension filter (e.g. '.pdf', '.py', '.txt')"
            },
            "count": {
                "type": "INTEGER",
                "description": "Number of results for largest or recent_files"
            },
            "depth": {
                "type": "INTEGER",
                "description": "Max depth level for tree view (default: 2)"
            }
        },
        "required": [
            "action"
        ]
    },
    "handler": file_controller,
}